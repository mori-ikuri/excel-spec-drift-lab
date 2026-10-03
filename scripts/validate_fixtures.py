#!/usr/bin/env python3
"""Validate this fixture contract, not infer missed changes from specifications.

Only openpyxl cell values and C# text lines are searched. A matching location
must have exactly one labels.csv or allowlist.csv entry with identical content.
The change-request document is deliberately excluded from this search.
"""

import argparse
import csv
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

from dump_cells import dump_version, serializable

ROOT = Path(__file__).resolve().parents[1]
PATTERN = re.compile(r"顧客名|CUSTOMER_NAME|CustomerName|cusNm|txtCusNm|(?<!\d)20(?!\d)")
LABEL_FIELDS = ["ID", "beforeファイル", "afterファイル", "シート名", "セル番地または行番号",
                "before内容", "after内容", "ラベル", "理由"]
ALLOW_FIELDS = ["版(before/after)", "ファイル", "シート名", "セル番地または行番号", "中身", "理由"]
BOOK_NAMES = {"DB定義書.xlsx", "画面設計書.xlsx", "帳票設計書.xlsx", "機能詳細設計書.xlsx"}
SOURCE_NAMES = {"Columns.cs", "Customer.cs", "CustomerForm.cs", "CustomerForm.Designer.cs"}
EXPECTED_LABELS = {
    "E1": "変更済み", "E2": "C", "E3": "変更済み", "E4": "A",
    "E5": "C", "E6": "A", "E7": "C", "E8": "A",
    "S1": "C", "S2": "B", "S3": "C", "S4": "A",
    "S5": "A", "S6": "C", "S7": "C", "S8": "C",
}


def read_csv(path, fields):
    with path.open(encoding="utf-8", newline="") as stream:
        reader = csv.DictReader(stream)
        if reader.fieldnames != fields:
            raise ValueError(f"{path.name}: incorrect CSV header")
        rows = list(reader)
    for number, row in enumerate(rows, start=2):
        if None in row or any(value is None for value in row.values()):
            raise ValueError(f"{path.name}:{number}: incorrect column count")
    return rows


def layout(ws):
    """Compare in-memory openpyxl formatting objects; never search XLSX bytes/XML."""
    return {
        "merges": sorted(str(area) for area in ws.merged_cells.ranges),
        "columns": [(name, d.width, d.hidden, d.min, d.max)
                    for name, d in sorted(ws.column_dimensions.items())],
        "rows": [(number, d.height, d.hidden) for number, d in sorted(ws.row_dimensions.items())],
        "cell_styles": [(cell.coordinate, tuple(cell._style or ()), cell.number_format)
                        for row in ws.iter_rows() for cell in row],
        "print_area": str(ws.print_area),
        "margins": str(ws.page_margins),
        "page_setup": str(ws.page_setup),
        "page_properties": str(ws.sheet_properties.pageSetUpPr),
        "print_options": str(ws.print_options),
        "row_breaks": [(b.id, b.min, b.max, b.man) for b in ws.row_breaks.brk],
        "column_breaks": [(b.id, b.min, b.max, b.man) for b in ws.col_breaks.brk],
        "sheet_view": str(ws.sheet_view),
    }


def validate(root):
    errors = []
    corpus, pairs, layouts = {}, {}, {}
    file_count = 0
    for version in ("before", "after"):
        for kind, names in (("docs", BOOK_NAMES), ("src", SOURCE_NAMES)):
            directory = root / "fixtures" / kind / version
            paths = sorted(p for p in directory.rglob("*") if p.is_file())
            if {p.relative_to(directory).as_posix() for p in paths} != names:
                errors.append(f"{kind}/{version}: expected exactly {sorted(names)}")
            for path in paths:
                file_count += 1
                relative = path.relative_to(root).as_posix()
                if kind == "src":
                    content = path.read_text(encoding="utf-8")
                    pairs[version, kind, path.name] = content
                    for number, line in enumerate(content.splitlines(), start=1):
                        corpus[version, relative, "", str(number)] = line
                elif path.suffix.lower() == ".xlsx":
                    book = load_workbook(path, data_only=False)
                    try:
                        if len(book.worksheets) not in (3, 4) or book.sheetnames[:2] != ["表紙", "改訂履歴"]:
                            errors.append(f"{relative}: expected cover, history and 1–2 body sheets")
                        values = {}
                        for ws in book.worksheets:
                            layouts[version, path.name, ws.title] = layout(ws)
                            for row in ws.iter_rows():
                                for cell in row:
                                    values[ws.title, cell.coordinate] = serializable(cell.value)
                                    if not isinstance(cell, MergedCell) and cell.value is not None:
                                        corpus[version, relative, ws.title, cell.coordinate] = str(serializable(cell.value))
                        pairs[version, kind, path.name] = values
                    finally:
                        book.close()

    labels = read_csv(root / "expected" / "labels.csv", LABEL_FIELDS)
    allowlist = read_csv(root / "expected" / "allowlist.csv", ALLOW_FIELDS)
    ids = [row["ID"] for row in labels]
    if Counter(ids) != Counter(EXPECTED_LABELS.keys()):
        errors.append("labels.csv: expected E1–E8 and S1–S8 exactly once")
    mapped = {}

    def register(key, content, reason, origin):
        if key in mapped:
            errors.append(f"Duplicate mapping: {key} ({mapped[key]}, {origin})")
        mapped[key] = origin
        if not reason.strip():
            errors.append(f"{origin}: missing reason")
        if key not in corpus:
            errors.append(f"{origin}: nonexistent or non-anchor location {key}")
        elif corpus[key] != content:
            errors.append(f"{origin}: content mismatch at {key}")
        if not PATTERN.search(content):
            errors.append(f"{origin}: mapping does not contain a searched term")

    for row in labels:
        if row["ラベル"] != EXPECTED_LABELS.get(row["ID"]):
            errors.append(f"{row['ID']}: incorrect label")
        for version in ("before", "after"):
            register((version, row[f"{version}ファイル"], row["シート名"], row["セル番地または行番号"]),
                     row[f"{version}内容"], row["理由"], row["ID"])
    for number, row in enumerate(allowlist, start=2):
        register((row["版(before/after)"], row["ファイル"], row["シート名"], row["セル番地または行番号"]),
                 row["中身"], row["理由"], f"allowlist.csv:{number}")

    matches = {key for key, value in corpus.items() if PATTERN.search(value)}
    for key in sorted(matches - mapped.keys()):
        errors.append(f"Unmapped match: {key}: {corpus[key]}")

    expected_differences = set()
    for row in labels:
        if row["ID"] in ("E1", "E3"):
            expected_differences.add((Path(row["beforeファイル"]).name, row["シート名"], row["セル番地または行番号"]))
            if row["before内容"].replace("20", "30") != row["after内容"]:
                errors.append(f"{row['ID']}: expected only the specified 20 to 30 change")
        elif row["before内容"] != row["after内容"]:
            errors.append(f"{row['ID']}: this evaluation location must be unchanged")

    differences = set()
    for filename in BOOK_NAMES:
        before = pairs.get(("before", "docs", filename), {})
        after = pairs.get(("after", "docs", filename), {})
        if before.keys() != after.keys():
            errors.append(f"{filename}: cell positions or sheet names differ")
        for key in before.keys() | after.keys():
            if before.get(key) != after.get(key):
                differences.add((filename, *key))
    if differences != expected_differences or len(differences) != 2:
        errors.append(f"Expected only E1/E3 cell differences; actual: {sorted(differences)}")
    for version, filename, sheet_name in layouts:
        if version == "before" and layouts[version, filename, sheet_name] != layouts.get(("after", filename, sheet_name)):
            errors.append(f"{filename}/{sheet_name}: formatting or print settings differ")
    for filename in SOURCE_NAMES:
        if pairs.get(("before", "src", filename)) != pairs.get(("after", "src", filename)):
            errors.append(f"{filename}: before and after sources differ")
    for version in ("before", "after"):
        expected = root / "expected" / f"cells-{version}.txt"
        if not expected.is_file() or expected.read_text(encoding="utf-8") != dump_version(root, version):
            errors.append(f"cells-{version}.txt: missing or stale dump")

    if errors:
        print("FAIL")
        for error in errors:
            print(f"- {error}")
        return False
    print("PASS")
    print(f"Files scanned: {file_count} (8 XLSX, 8 C#; change-request.md excluded)")
    print(f"Evaluation locations: {len(labels)} (32 before/after mappings)")
    print(f"Allowlist locations: {len(allowlist)}")
    for version in ("before", "after"):
        print(f"Matched locations ({version}): {sum(key[0] == version for key in matches)}")
    print("Unmapped locations: 0")
    print("Cell differences: E1, E3 only; layouts and all C# sources identical")
    print("Saved cell dumps: match")
    return True


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=ROOT)
    args = parser.parse_args()
    try:
        success = validate(args.root.resolve())
    except (OSError, ValueError, KeyError) as error:
        print(f"FAIL: {error}")
        success = False
    raise SystemExit(0 if success else 1)


if __name__ == "__main__":
    main()
