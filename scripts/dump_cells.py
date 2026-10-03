#!/usr/bin/env python3
"""Dump every cell in each worksheet's used rectangle, including blank cells.

Each line is a JSON array [filename, sheet name, cell address, value].
Filenames are version-relative so before/after dumps differ only in cell values.
Dates are ISO 8601 strings and blank/merged non-anchor cells are JSON null.
No timestamps or binary XLSX hashes are used.
"""

import argparse
import json
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]


def serializable(value):
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    return value


def dump_version(root, version):
    lines = []
    directory = root / "fixtures" / "docs" / version
    for path in sorted(directory.rglob("*.xlsx")):
        book = load_workbook(path, data_only=False)
        try:
            for ws in book.worksheets:
                for row in ws.iter_rows():
                    for cell in row:
                        lines.append(json.dumps([
                            path.relative_to(directory).as_posix(), ws.title,
                            cell.coordinate, serializable(cell.value),
                        ], ensure_ascii=False, separators=(",", ":")))
        finally:
            book.close()
    return "\n".join(lines) + "\n"


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=ROOT)
    args = parser.parse_args()
    output = args.root / "expected"
    output.mkdir(parents=True, exist_ok=True)
    for version in ("before", "after"):
        text = dump_version(args.root, version)
        path = output / f"cells-{version}.txt"
        path.write_text(text, encoding="utf-8", newline="\n")
        print(f"{path.name}: {len(text.splitlines())} cells")


if __name__ == "__main__":
    main()
