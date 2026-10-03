#!/usr/bin/env python3
"""Generate the fictional documents and sources using fixed content/layout values."""

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.properties import PageSetupProperties

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "fixtures"
# Fine grid: approximately 15.5 px wide and 15 px high at standard display DPI.
COLUMN_WIDTH = 1.5
ROW_HEIGHT = 11.25
LAST_COLUMN = 48
LAST_ROW = 48
FONT_NAME = "Noto Sans CJK JP"
FONT_SIZE = 10
TITLE_SIZE = 14
PRINT_SCALE = 90
PRINT_AREA = "A1:AV48"
MARGIN = 0.3
FIXED_DATE = datetime(2026, 10, 3)
GRID = Side(style="hair", color="D9D9D9")
RULE = Side(style="thin", color="595959")
HEADER_FILL = "E7E6E6"


def sheet(book, name):
    ws = book.create_sheet(name)
    ws.sheet_view.showGridLines = True
    ws.sheet_view.zoomScale = 100
    for column in range(1, LAST_COLUMN + 1):
        ws.column_dimensions[get_column_letter(column)].width = COLUMN_WIDTH
    for row in range(1, LAST_ROW + 1):
        ws.row_dimensions[row].height = ROW_HEIGHT
        for column in range(1, LAST_COLUMN + 1):
            cell = ws.cell(row, column)
            cell.font = Font(name=FONT_NAME, size=FONT_SIZE, color="222222")
            cell.alignment = Alignment(vertical="center")
            cell.border = Border(left=GRID, right=GRID, top=GRID, bottom=GRID)
    ws.print_area = PRINT_AREA
    ws.print_options.horizontalCentered = True
    ws.page_margins = PageMargins(left=MARGIN, right=MARGIN, top=MARGIN,
                                 bottom=MARGIN, header=0.1, footer=0.1)
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=False)
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.page_setup.scale = PRINT_SCALE
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    # Manual boundaries are at the edge of the single-page print area.
    ws.row_breaks.append(Break(id=LAST_ROW))
    ws.col_breaks.append(Break(id=LAST_COLUMN))
    return ws


def block(ws, area, value, *, header=False, title=False):
    """Put a value at the top-left of a merged grid block."""
    ws.merge_cells(area)
    cells = ws[area]
    for row in cells:
        for cell in row:
            cell.border = Border(left=RULE, right=RULE, top=RULE, bottom=RULE)
            if header:
                cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
    cell = cells[0][0]
    cell.value = value
    cell.font = Font(name=FONT_NAME, size=TITLE_SIZE if title else FONT_SIZE,
                     bold=header or title, color="222222")
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True,
                               indent=1)
    if isinstance(value, datetime):
        cell.number_format = "yyyy-mm-dd"


def book_with_front_matter(title, document_id):
    book = Workbook()
    book.remove(book.active)
    book.properties.creator = ""
    book.properties.lastModifiedBy = ""
    book.properties.created = FIXED_DATE
    book.properties.modified = FIXED_DATE
    book.properties.title = title
    cover = sheet(book, "表紙")
    block(cover, "C5:AT8", "顧客管理システム", title=True)
    block(cover, "C11:AT14", title, title=True)
    block(cover, "C19:N21", "文書番号", header=True)
    block(cover, "O19:AT21", document_id)
    block(cover, "C24:N26", "作成日", header=True)
    block(cover, "O24:AT26", FIXED_DATE)
    block(cover, "C29:N31", "版数", header=True)
    block(cover, "O29:AT31", "1.0")
    block(cover, "C36:AT39", "架空のダミーシステム")
    history = sheet(book, "改訂履歴")
    block(history, "C3:AT6", "改訂履歴", title=True)
    block(history, "C10:J12", "版数", header=True)
    block(history, "K10:V12", "日付", header=True)
    block(history, "W10:AT12", "内容", header=True)
    block(history, "C13:J15", "1.0")
    block(history, "K13:V15", FIXED_DATE)
    block(history, "W13:AT15", "初版作成")
    return book


def document_books(length):
    db = book_with_front_matter("DB定義書", "DB-001")
    ws = sheet(db, "テーブル定義")
    block(ws, "C3:AT6", "顧客テーブル", title=True)
    block(ws, "C9:H11", "No.", header=True)
    block(ws, "I9:AT11", "項目名 / 列名 / データ型", header=True)
    for row, number, value in [
        (12, 1, "顧客ID / CUSTOMER_ID / INT"),
        (16, 2, f"顧客名 / CUSTOMER_NAME / NVARCHAR({length})"),
        (20, 3, "顧客名カナ / CUSTOMER_NAME_KANA / NVARCHAR(20)"),
    ]:
        block(ws, f"C{row}:H{row + 3}", number)
        block(ws, f"I{row}:AT{row + 3}", value)

    screen = book_with_front_matter("画面設計書", "UI-001")
    ws = sheet(screen, "画面仕様")
    block(ws, "C3:AT6", "顧客照会画面", title=True)
    block(ws, "C9:AT11", "項目一覧", header=True)
    block(ws, "C12:AT15", f"顧客名・最大{length}桁")
    block(ws, "C18:AT20", "入力チェック", header=True)
    block(ws, "C21:AT24", "顧客名は20文字以内")
    block(ws, "C27:AT29", "データ取得SQL", header=True)
    block(ws, "C30:AT34", "SELECT CUSTOMER_NAME FROM CUSTOMER WHERE CUSTOMER_ID = @id;")

    report = book_with_front_matter("帳票設計書", "RP-001")
    ws = sheet(report, "帳票仕様")
    block(ws, "C3:AT6", "顧客一覧表", title=True)
    block(ws, "C9:AT11", "印字項目", header=True)
    block(ws, "C12:AT15", "顧客名・印字桁20")
    block(ws, "C18:AT20", "見出し", header=True)
    block(ws, "C21:AT24", "顧客名")

    detail = book_with_front_matter("機能詳細設計書", "FN-001")
    ws = sheet(detail, "処理仕様")
    block(ws, "C3:AT6", "印字データ編集", title=True)
    block(ws, "C9:AT11", "処理内容", header=True)
    block(ws, "C12:AT15", "顧客名を20文字で切り詰める")
    return {"DB定義書.xlsx": db, "画面設計書.xlsx": screen,
            "帳票設計書.xlsx": report, "機能詳細設計書.xlsx": detail}


SOURCES = {
    "Columns.cs": '''namespace CustomerManagement
{
    internal static class Columns
    {
        public const string CUSTOMER_ID = "CUSTOMER_ID";
        public const string CUSTOMER_NAME = "CUSTOMER_NAME";
        public const string CUSTOMER_NAME_KANA = "CUSTOMER_NAME_KANA";
    }
}
''',
    "Customer.cs": '''namespace CustomerManagement
{
    public class Customer
    {
        public int CustomerId { get; set; }
        public string CustomerName { get; set; } // 顧客名(20桁)
    }
}
''',
    "CustomerForm.cs": '''using System;
using System.Collections.Generic;
using System.Windows.Forms;

namespace CustomerManagement
{
    public partial class CustomerForm : Form
    {
        public CustomerForm()
        {
            InitializeComponent();
        }

        public void ShowCustomer(IDictionary<string, object> data)
        {
            var cusNm = data[Columns.CUSTOMER_NAME].ToString();
            if (cusNm.Length > 20)
            {
                throw new ArgumentException("入力値が長すぎます。");
            }
            txtCusNm.Text = cusNm;
        }
    }
}
''',
    "CustomerForm.Designer.cs": '''using System.Drawing;
using System.Windows.Forms;

namespace CustomerManagement
{
    public partial class CustomerForm
    {
        private TextBox txtCusNm = new TextBox
        {
            Location = new Point(24, 24),
            Size = new Size(264, 24),
            TabIndex = 0
        };

        private TextBox txtCusNmKana = new TextBox
        {
            Location = new Point(24, 64),
            Size = new Size(264, 24),
            TabIndex = 1
        };

        private void InitializeComponent()
        {
            SuspendLayout();
            txtCusNm.MaxLength = 20;
            txtCusNmKana.MaxLength = 20;
            Controls.AddRange(new Control[] { txtCusNm, txtCusNmKana });
            ClientSize = new Size(312, 112);
            Text = "顧客情報";
            ResumeLayout(false);
            PerformLayout();
        }
    }
}
''',
}


def main():
    FIXTURES.mkdir(exist_ok=True)
    (FIXTURES / "change-request.md").write_text(
        "顧客名(CUSTOMER_NAME)の桁数を20桁から30桁に変更する。理由:20では足りず、"
        "25でもまた足りなくなる恐れがあるため、余裕を見て30とする\n", encoding="utf-8")
    for version, length in (("before", 20), ("after", 30)):
        docs = FIXTURES / "docs" / version
        src = FIXTURES / "src" / version
        docs.mkdir(parents=True, exist_ok=True)
        src.mkdir(parents=True, exist_ok=True)
        for filename, book in document_books(length).items():
            book.save(docs / filename)
            book.close()
        for filename, source in SOURCES.items():
            (src / filename).write_text(source, encoding="utf-8", newline="\n")
    print("Generated: 8 workbooks, 8 C# files, 1 change request")


if __name__ == "__main__":
    main()
